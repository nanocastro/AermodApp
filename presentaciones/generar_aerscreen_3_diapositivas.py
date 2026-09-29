from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path(__file__).with_name("AERSCREEN_4_diapositivas.pptx")

NAVY = RGBColor(13, 36, 54)
TEAL = RGBColor(0, 126, 133)
CYAN = RGBColor(52, 188, 205)
SKY = RGBColor(218, 243, 246)
PALE = RGBColor(244, 249, 249)
WHITE = RGBColor(255, 255, 255)
INK = RGBColor(28, 45, 56)
MUTED = RGBColor(91, 111, 121)
AMBER = RGBColor(244, 169, 49)
GREEN = RGBColor(70, 156, 110)
RED = RGBColor(208, 82, 73)
LINE = RGBColor(194, 214, 218)

FONT = "Aptos"
SOURCE = "Fuente: U.S. EPA, AERSCREEN User’s Guide (EPA-454/B-21-005, 2021) · epa.gov/scram"


def rect(slide, x, y, w, h, fill, radius=True, line=None, line_width=1):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    shape.line.width = Pt(line_width)
    if radius:
        try:
            shape.adjustments[0] = 0.12
        except (IndexError, ValueError):
            pass
    return shape


def text(slide, value, x, y, w, h, size=18, color=INK, bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.05,
         font=FONT, line_spacing=1.0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    p = frame.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def bullets(slide, items, x, y, w, h, size=15, color=INK, bullet_color=None, spacing=6):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.02)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    for idx, item in enumerate(items):
        p = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        p.text = f"•  {item}"
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(spacing)
        p.line_spacing = 1.0
    return box


def title(slide, number, heading, subheading=None):
    text(slide, f"0{number}", 0.55, 0.35, 0.55, 0.35, 14, TEAL, True)
    text(slide, heading, 1.13, 0.25, 11.4, 0.55, 28, NAVY, True)
    if subheading:
        text(slide, subheading, 1.14, 0.8, 11.2, 0.4, 12.5, MUTED)
    rect(slide, 0.55, 1.15, 12.2, 0.02, TEAL, False)


def footer(slide):
    text(slide, SOURCE, 0.55, 7.18, 9.8, 0.18, 8, MUTED)
    text(slide, "AERSCREEN · síntesis conceptual", 10.55, 7.18, 2.2, 0.18, 8, MUTED, align=PP_ALIGN.RIGHT)


def pill(slide, label, x, y, w, fill, color):
    rect(slide, x, y, w, 0.42, fill, True)
    text(slide, label, x + 0.08, y + 0.02, w - 0.16, 0.34, 11, color, True,
         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def arrow(slide, x, y, w, h=0.32, color=TEAL):
    sh = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.color.rgb = color
    return sh


def circle(slide, x, y, d, fill, line=None):
    sh = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line or fill
    return sh


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

# Diapositiva 1 — definición
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = PALE

text(slide, "AERSCREEN", 0.68, 0.52, 7.5, 0.85, 42, NAVY, True)
text(slide, "Una primera mirada conservadora antes del modelado refinado", 0.72, 1.33, 7.3, 0.6, 22, TEAL, True)
text(slide,
     "Es el modelo de screening recomendado por la EPA y basado en AERMOD. Busca la combinación meteorológica y la distancia que generan el mayor impacto esperado de una fuente, sin exigir una serie meteorológica horaria completa.",
     0.72, 2.12, 6.55, 1.65, 18, INK, False, line_spacing=1.08)

pill(slide, "RECOMENDADO POR EPA", 0.72, 4.05, 1.95, SKY, TEAL)
pill(slide, "UNA FUENTE", 2.82, 4.05, 1.55, RGBColor(232, 240, 244), NAVY)
pill(slide, "PEOR CASO", 4.52, 4.05, 1.55, RGBColor(255, 239, 210), RGBColor(154, 97, 10))

rect(slide, 0.72, 4.72, 6.5, 1.32, NAVY, True)
text(slide, "¿Qué entrega?", 0.98, 4.92, 1.7, 0.3, 14, CYAN, True)
text(slide, "Máxima concentración de 1 hora y estimaciones conservadoras para 3 h, 8 h, 24 h y anual.",
     0.98, 5.26, 5.8, 0.55, 17, WHITE, True)

# Gráfico conceptual fuente-pluma-receptor
rect(slide, 7.75, 0.48, 4.9, 5.95, WHITE, True, LINE)
text(slide, "SCREENING ≠ PRONÓSTICO HORARIO", 8.08, 0.82, 4.25, 0.4, 12, TEAL, True, align=PP_ALIGN.CENTER)
rect(slide, 8.03, 4.76, 4.35, 0.08, RGBColor(164, 178, 181), False)
rect(slide, 8.52, 2.73, 0.44, 2.03, NAVY, False)
rect(slide, 8.43, 2.58, 0.62, 0.18, NAVY, False)
circle(slide, 8.56, 2.28, 0.35, AMBER)
text(slide, "fuente", 8.18, 4.96, 1.05, 0.3, 11, MUTED, True, align=PP_ALIGN.CENTER)

# Pluma como bandas transparentes aproximadas
for i, (x, y, w, h, c) in enumerate([
    (8.87, 2.40, 1.05, 0.42, RGBColor(145, 220, 225)),
    (9.50, 2.51, 1.28, 0.62, RGBColor(181, 231, 234)),
    (10.35, 2.69, 1.47, 0.88, RGBColor(215, 242, 244)),
]):
    sh = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = c; sh.line.color.rgb = c

circle(slide, 11.45, 4.43, 0.38, RED)
text(slide, "receptor crítico", 10.78, 4.96, 1.75, 0.3, 11, MUTED, True, align=PP_ALIGN.CENTER)
arrow(slide, 9.20, 5.58, 2.15, 0.28, MUTED)
text(slide, "explora distancias y condiciones", 8.95, 5.90, 2.85, 0.34, 10.5, MUTED, align=PP_ALIGN.CENTER)

text(slide, "Resultado deliberadamente conservador: sirve para decidir si hace falta un análisis refinado.",
     0.72, 6.42, 11.9, 0.5, 15, NAVY, True, align=PP_ALIGN.CENTER)
footer(slide)

# Diapositiva 2 — arquitectura
slide = prs.slides.add_slide(blank)
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = WHITE
title(slide, 2, "¿Qué modelos lo componen y cómo interactúan?",
      "AERSCREEN organiza el flujo; cada programa resuelve una parte distinta del problema.")

# Entrada
rect(slide, 0.55, 2.40, 1.72, 1.42, NAVY, True)
text(slide, "DATOS DEL\nESCENARIO", 0.66, 2.67, 1.50, 0.62, 13.5, WHITE, True,
     align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
text(slide, "fuente · entorno", 0.75, 3.35, 1.32, 0.2, 9.5, CYAN, align=PP_ALIGN.CENTER)
arrow(slide, 2.30, 2.93, 0.55, 0.26, TEAL)

# Orquestador
rect(slide, 2.87, 2.15, 2.08, 1.92, SKY, True, TEAL, 2)
text(slide, "AERSCREEN", 3.01, 2.54, 1.80, 0.36, 16.5, NAVY, True, align=PP_ALIGN.CENTER)
text(slide, "valida · coordina\nbusca · refina", 3.17, 3.02, 1.48, 0.58, 11.5, TEAL,
     align=PP_ALIGN.CENTER)

# Preprocesadores
cards = [
    ("MAKEMET", "matriz meteorológica\nsintética", 1.33, AMBER),
    ("AERMAP*", "cotas de fuente\ny receptores", 2.75, GREEN),
    ("BPIPPRM*", "parámetros PRIME\nde edificios", 4.17, RED),
]
for label, desc, y, accent in cards:
    rect(slide, 5.56, y, 2.15, 1.02, PALE, True, accent, 1.5)
    circle(slide, 5.76, y + 0.30, 0.38, accent)
    text(slide, label, 6.23, y + 0.16, 1.25, 0.25, 14, NAVY, True)
    text(slide, desc, 6.23, y + 0.46, 1.25, 0.40, 9.8, MUTED)
    arrow(slide, 5.02, y + 0.37, 0.47, 0.22, RGBColor(153, 179, 183))
    arrow(slide, 7.78, y + 0.37, 0.48, 0.22, accent)

# AERMOD y resultados
rect(slide, 8.30, 2.14, 2.05, 1.94, NAVY, True)
text(slide, "AERMOD", 8.55, 2.49, 1.55, 0.36, 20, WHITE, True, align=PP_ALIGN.CENTER)
pill(slide, "MODELOPT SCREEN", 8.54, 2.98, 1.57, RGBColor(34, 72, 91), CYAN)
text(slide, "motor de dispersión", 8.56, 3.53, 1.52, 0.25, 10.5, WHITE, align=PP_ALIGN.CENTER)
arrow(slide, 10.42, 2.93, 0.55, 0.26, TEAL)
rect(slide, 11.00, 2.40, 1.78, 1.42, SKY, True, TEAL, 1.5)
text(slide, "MÁXIMO\nCONSERVADOR", 11.12, 2.69, 1.54, 0.58, 11.5, NAVY, True,
     align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
text(slide, "1 h + factores", 11.23, 3.36, 1.32, 0.2, 9.5, TEAL, align=PP_ALIGN.CENTER)

text(slide, "* Se activan cuando el escenario incluye terreno y/o downwash.", 5.55, 5.43, 4.55, 0.32, 11, MUTED)
rect(slide, 1.05, 6.05, 11.22, 0.63, PALE, True, LINE)
text(slide,
     "MAKEMET genera la meteorología → AERMAP aporta elevaciones → BPIPPRM caracteriza edificios → AERMOD calcula → AERSCREEN conserva el peor caso.",
     1.25, 6.18, 10.82, 0.34, 13.5, NAVY, True, align=PP_ALIGN.CENTER)
footer(slide)

# Diapositiva 3 — entradas
slide = prs.slides.add_slide(blank)
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = PALE
title(slide, 3, "¿Qué datos utiliza como entrada?",
      "Los datos describen la emisión, el ambiente y la geometría que condiciona la dispersión.")

columns = [
    (0.55, "1", "FUENTE EMISORA", TEAL,
     ["Tipo de fuente y tasa de emisión", "Altura y diámetro de chimenea", "Temperatura y velocidad de salida", "Entorno rural o urbano", "Ubicación y cota de base"]),
    (4.55, "2", "METEOROLOGÍA Y SUPERFICIE", AMBER,
     ["Temperaturas ambiente mínima y máxima", "Albedo", "Razón de Bowen", "Longitud de rugosidad", "Variación por período y sector"]),
    (8.55, "3", "GEOMETRÍA DEL ENTORNO", GREEN,
     ["Terreno plano, elevado o complejo", "DEM y coordenadas, si aplica", "Edificios: altura, planta y orientación", "Rango de búsqueda y receptores discretos", "Altura de receptor opcional"]),
]

for x, num, heading, accent, items in columns:
    rect(slide, x, 1.48, 3.68, 4.88, WHITE, True, LINE)
    circle(slide, x + 0.27, 1.77, 0.58, accent)
    text(slide, num, x + 0.27, 1.84, 0.58, 0.35, 17, WHITE, True,
         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    text(slide, heading, x + 1.00, 1.77, 2.34, 0.58, 14, NAVY, True,
         valign=MSO_ANCHOR.MIDDLE)
    rect(slide, x + 0.27, 2.54, 3.14, 0.04, accent, False)
    bullets(slide, items, x + 0.30, 2.83, 3.03, 2.65, 14, INK, spacing=10)
    if num == "2":
        pill(slide, "NO REQUIERE SERIE HORARIA", x + 0.52, 5.66, 2.64, RGBColor(255, 243, 219), RGBColor(145, 91, 8))
    elif num == "1":
        pill(slide, "CASO TÍPICO: FUENTE PUNTUAL", x + 0.47, 5.66, 2.74, SKY, TEAL)
    else:
        pill(slide, "AERMAP / BPIPPRM SI APLICA", x + 0.52, 5.66, 2.64, RGBColor(226, 242, 233), GREEN)

text(slide,
     "Idea clave: el resultado es tan defendible como los supuestos de entrada. Conservador no significa automáticamente representativo ni regulatorio.",
     0.75, 6.58, 11.85, 0.45, 14, NAVY, True, align=PP_ALIGN.CENTER)
footer(slide)

# Diapositiva 4 — resultados, uso y límites
slide = prs.slides.add_slide(blank)
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = WHITE
title(slide, 4, "¿Qué resultados entrega y cómo se interpretan?",
      "El valor del screening está en orientar decisiones tempranas, no en reemplazar un estudio refinado.")

# Tres resultados principales
result_cards = [
    (0.65, "01", "MÁXIMO DE 1 HORA", "Concentración máxima estimada y distancia crítica desde la fuente.", TEAL),
    (4.55, "02", "OTROS PERÍODOS", "Estimaciones conservadoras de 3 h, 8 h, 24 h y anual mediante factores.", AMBER),
    (8.45, "03", "CONDICIÓN CRÍTICA", "Combinación meteorológica y geometría que controla el peor caso.", GREEN),
]
for x, num, heading, body, accent in result_cards:
    rect(slide, x, 1.48, 3.58, 1.55, PALE, True, LINE)
    circle(slide, x + 0.22, 1.71, 0.52, accent)
    text(slide, num, x + 0.22, 1.77, 0.52, 0.30, 13, WHITE, True,
         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    text(slide, heading, x + 0.88, 1.66, 2.45, 0.30, 13.5, NAVY, True)
    text(slide, body, x + 0.88, 2.04, 2.42, 0.72, 11.8, MUTED, line_spacing=1.0)

# Panel de utilidad
rect(slide, 0.65, 3.42, 5.82, 2.70, RGBColor(230, 245, 238), True, GREEN, 1.5)
text(slide, "¿PARA QUÉ SIRVE?", 0.98, 3.74, 2.55, 0.36, 16, GREEN, True)
bullets(slide, [
    "Identificar rápidamente el escenario de mayor impacto.",
    "Comparar alternativas de ubicación, altura o emisión.",
    "Priorizar controles y datos que requieren mayor detalle.",
    "Decidir si corresponde avanzar a una modelación refinada.",
], 0.98, 4.20, 5.00, 1.62, 12.2, NAVY, spacing=3)

# Panel de límites
rect(slide, 6.82, 3.42, 5.86, 2.70, RGBColor(255, 244, 224), True, AMBER, 1.5)
text(slide, "LÍMITES DE INTERPRETACIÓN", 7.15, 3.74, 3.45, 0.36, 16, RGBColor(145, 91, 8), True)
bullets(slide, [
    "La meteorología es sintética: no representa una cronología real.",
    "El grado de conservadurismo varía según el caso analizado.",
    "No describe cuándo ocurrirá el máximo ni su frecuencia.",
    "No sustituye la revisión técnica, científica o normativa.",
], 7.15, 4.20, 5.05, 1.62, 12.2, NAVY, spacing=3)

rect(slide, 1.32, 6.42, 10.70, 0.54, NAVY, True)
text(slide, "LECTURA CORRECTA: señal de alerta conservadora → base para decidir el siguiente nivel de análisis",
     1.55, 6.52, 10.24, 0.28, 13.5, WHITE, True, align=PP_ALIGN.CENTER)
footer(slide)

prs.core_properties.title = "AERSCREEN — definición, componentes, entradas y resultados"
prs.core_properties.subject = "Presentación técnica introductoria de cuatro diapositivas"
prs.core_properties.author = "Proyecto AERMOD"
prs.core_properties.keywords = "AERSCREEN, AERMOD, MAKEMET, AERMAP, BPIPPRM"
prs.save(OUT)
print(OUT)
